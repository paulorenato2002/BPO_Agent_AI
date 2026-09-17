"""Listas de valores fora do PDF: planilhas de VT/VA, arquivos .txt e dados
colados na mensagem.

Cada lista vira um `Documento` do tipo "lista", com um item por pessoa (ou
favorecido) e o rótulo do benefício (VT, VA, VR) quando dá para saber. A
leitura é por regra: coluna de nome, coluna de valor e linha de total. O que
não se encaixa vira alerta, nunca valor inventado.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from hashlib import sha256

from core import Documento, Item, norm

# Rótulo → palavras que o identificam (já normalizadas). Siglas só valem como palavra inteira.
BENEFICIOS = {
    "VT": ("VALE TRANSPORTE", "VALES TRANSPORTE", "TRANSPORTE", "VT", "PASSAGEM", "PASSAGENS",
           "MOBILIDADE", "RIOCARD", "BILHETE UNICO", "SPTRANS", "SETPS", "SETRANSP", "DFTRANS"),
    "VA": ("VALE ALIMENTACAO", "ALIMENTACAO", "VA", "CESTA BASICA"),
    "VR": ("VALE REFEICAO", "REFEICAO", "VR"),
}
# Operadoras que emitem mais de um benefício: ajudam a achar o lançamento, não o rótulo.
OPERADORAS = ("ALELO", "SODEXO", "PLUXEE", "TICKET", "VR BENEFICIOS", "FLASH", "CAJU", "IFOOD BENEFICIOS",
              "SWILE", "UP BRASIL", "VEROCARD", "GOODCARD", "BEN VISA", "VALECARD")

COLUNAS_NOME = ("NOME", "FUNCIONARIO", "COLABORADOR", "BENEFICIARIO", "EMPREGADO", "FAVORECIDO", "SERVIDOR",
                "TRABALHADOR", "PESSOA", "FORNECEDOR")
COLUNAS_CPF = ("CPF",)
RE_VALOR_TEXTO = re.compile(r"(?:R\$\s*)?(-?\d{1,3}(?:\.\d{3})+,\d{2}|-?\d+,\d{2}|-?\d+\.\d{2}(?!\d))")
RE_CPF = re.compile(r"\d{3}\.?\d{3}\.?\d{3}-?\d{2}")
LIMITE_LINHAS = 2000


def _palavras(texto: str) -> str:
    return f" {norm(texto)} "


def rotulo_de(*textos: str) -> str:
    """VT, VA ou VR, quando algum dos textos deixa isso claro; senão vazio."""
    achados = []
    for texto in textos:
        t = _palavras(texto)
        for rotulo, chaves in BENEFICIOS.items():
            if any(f" {c} " in t for c in chaves):
                achados.append(rotulo)
    unicos = list(dict.fromkeys(achados))
    return unicos[0] if len(unicos) == 1 else ""


def fala_de(rotulo: str, item: Item) -> bool:
    """O lançamento do contas a pagar é deste benefício? Categoria e descrição
    dizem o benefício; do fornecedor só vale o nome de operadora."""
    if rotulo not in BENEFICIOS:
        return False
    t = _palavras(item.categoria + " " + item.descricao)
    if any(f" {c} " in t for c in BENEFICIOS[rotulo]):
        return True
    f = _palavras(item.nome)
    return any(f" {o} " in f for o in OPERADORAS) and not any(
        f" {c} " in t for r, cs in BENEFICIOS.items() if r != rotulo for c in cs)


def _centavos(valor) -> int | None:
    if isinstance(valor, dict):                      # célula com fórmula vinda do Excel
        valor = valor.get("result", valor.get("text"))
    if isinstance(valor, bool) or valor is None:
        return None
    if isinstance(valor, (int, float, Decimal)):
        try:
            return int((Decimal(str(valor)) * 100).quantize(Decimal("1")))
        except InvalidOperation:
            return None
    texto = str(valor).strip()
    if not texto:
        return None
    m = RE_VALOR_TEXTO.fullmatch(texto.replace("R$", "").strip())
    if not m:
        return None
    bruto = m.group(1)
    if "," in bruto:
        bruto = bruto.replace(".", "").replace(",", ".")
    return int((Decimal(bruto) * 100).quantize(Decimal("1")))


def _texto(valor) -> str:
    if isinstance(valor, dict):
        valor = valor.get("result") or valor.get("text") or " ".join(
            p.get("text", "") for p in valor.get("richText", []) if isinstance(p, dict))
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    return "" if valor is None else str(valor).strip()


def _pontos_valor(cabecalho: str) -> int:
    h = norm(cabecalho)
    if not h:
        return -99
    palavras = set(h.split())
    pontos = 0
    if palavras & {"TOTAL", "PAGAR", "CREDITAR", "CREDITO", "RECARGA", "LIQUIDO", "DEPOSITO", "PEDIDO"}:
        pontos += 3
    if palavras & {"VALOR", "VLR", "VL", "R"}:
        pontos += 1
    if rotulo_de(cabecalho):
        pontos += 2
    if palavras & {"DIA", "DIARIO", "DIARIA", "UNITARIO", "UNIT", "TARIFA", "DIAS", "QTD", "QUANTIDADE",
                   "DESCONTO", "PERCENTUAL", "CPF", "MATRICULA", "CODIGO", "SALARIO", "BASE"}:
        pontos -= 5
    return pontos


def _eh_total(texto: str) -> bool:
    return bool(re.match(r"^(TOTAL|TOTAIS|SOMA|VALOR TOTAL)\b", norm(texto)))


def _achar_cabecalho(linhas: list[list]) -> int | None:
    for i, linha in enumerate(linhas[:30]):
        textos = [norm(_texto(c)) for c in linha]
        if any(any(t == n or t.startswith(n + " ") or t.endswith(" " + n) for n in COLUNAS_NOME) for t in textos):
            return i
    return None


def ler_tabela(nome: str, linhas: list[list], aba: str = "") -> list[Documento]:
    """Uma tabela (lista de linhas) vira uma lista por coluna de valor relevante."""
    hash_ = sha256(repr((nome, aba, linhas[:LIMITE_LINHAS])).encode()).hexdigest()
    base = Documento(nome if not aba else f"{nome} — {aba}", hash_, "lista", 1)
    i_cab = _achar_cabecalho(linhas)
    if i_cab is None:
        base.alertas.append("Não encontrei a coluna de nome (NOME, FUNCIONÁRIO, COLABORADOR...).")
        return [base]
    cab = [_texto(c) for c in linhas[i_cab]]
    col_nome = next(j for j, c in enumerate(cab)
                    if any(norm(c) == n or norm(c).startswith(n + " ") or norm(c).endswith(" " + n) for n in COLUNAS_NOME))
    col_cpf = next((j for j, c in enumerate(cab) if norm(c) in COLUNAS_CPF or norm(c).startswith("CPF")), None)
    corpo = linhas[i_cab + 1:i_cab + 1 + LIMITE_LINHAS]

    # Uma coluna por benefício (NOME | VT | VA) vira uma lista para cada.
    por_beneficio = [(j, rotulo_de(c)) for j, c in enumerate(cab) if j != col_nome and rotulo_de(c)
                     and _pontos_valor(c) > 0 and any(_centavos(l[j]) is not None for l in corpo if j < len(l))]
    if len({r for _, r in por_beneficio}) < 2:
        candidatas = [(j, _pontos_valor(c)) for j, c in enumerate(cab) if j not in (col_nome, col_cpf)]
        candidatas = [(j, p) for j, p in candidatas
                      if p > -5 and any(_centavos(l[j]) is not None for l in corpo if j < len(l))]
        if not candidatas:
            base.alertas.append("Não encontrei a coluna de valor.")
            return [base]
        j_valor = max(candidatas, key=lambda x: (x[1], x[0]))[0]
        rotulo = rotulo_de(cab[j_valor]) or rotulo_de(nome, aba, " ".join(_texto(c) for l in linhas[:i_cab] for c in l))
        por_beneficio = [(j_valor, rotulo)]

    docs = []
    for j, rotulo in por_beneficio:
        doc = Documento(base.arquivo + (f" ({rotulo})" if len(por_beneficio) > 1 else ""), f"{hash_}:{j}",
                        "lista", 1, rotulo=rotulo)
        for n, linha in enumerate(corpo, i_cab + 2):
            celulas = list(linha) + [None] * (max(col_nome, j) + 1 - len(linha))
            quem, valor = _texto(celulas[col_nome]), _centavos(celulas[j])
            if _eh_total(quem) or (not quem and valor and _eh_total(" ".join(_texto(c) for c in linha))):
                doc.total_impresso = valor
                continue
            if not quem:
                continue
            if valor is None:
                if _texto(celulas[j]):
                    doc.alertas.append(f"Linha {n}: {quem} com valor ilegível em \"{cab[j]}\" ({_texto(celulas[j])}).")
                continue
            if valor == 0:
                continue
            cpf = _texto(celulas[col_cpf]) if col_cpf is not None and col_cpf < len(celulas) else ""
            doc.itens.append(Item(str(len(doc.itens) + 1), quem, valor, descricao=rotulo or "lista",
                                  documento=cpf, pagina=1, evidencia=" | ".join(_texto(c) for c in linha)))
        if doc.total_impresso is None:
            doc.total_impresso = doc.total      # sem linha de total, não há o que comparar
        elif doc.total_impresso != doc.total:
            doc.alertas.append("A soma das linhas não bate com o total da planilha.")
        if not doc.itens:
            doc.alertas.append("Nenhuma linha com nome e valor.")
        docs.append(doc)
    return docs


def ler_texto(nome: str, texto: str) -> list[Documento]:
    """Linhas "Nome — R$ 150,00". Linha sem valor que cita VT/VA abre uma seção."""
    listas: dict[str, Documento] = {}
    atual = rotulo_de(nome)
    avisos = []
    for n, bruta in enumerate((texto or "").splitlines()[:LIMITE_LINHAS], 1):
        linha = bruta.strip()
        if not linha:
            continue
        valores = list(RE_VALOR_TEXTO.finditer(linha))
        if not valores:
            atual = rotulo_de(linha) or atual
            continue
        m = valores[-1]
        resto = (linha[:m.start()] + " " + linha[m.end():]).replace("R$", " ")
        resto = RE_CPF.sub(" ", resto)
        rotulo = rotulo_de(resto) or atual
        quem = re.sub(r"\s+", " ", re.sub(r"[-–—:;|=\t]+", " ", resto)).strip()
        # O rótulo no começo da linha ("VT Fulano") não faz parte do nome.
        for chave in BENEFICIOS.get(rotulo, ()):
            quem = re.sub(rf"^(?:{chave})\b\s*", "", quem, flags=re.I)
        quem = quem.strip(" .,")
        valor = _centavos(m.group(1))
        doc = listas.setdefault(rotulo, Documento(nome + (f" ({rotulo})" if rotulo else ""), "", "lista", 1,
                                                  rotulo=rotulo))
        if _eh_total(quem or linha):
            doc.total_impresso = valor
            continue
        if not quem or len(norm(quem)) < 2:
            avisos.append(f"Linha {n} sem nome: \"{linha}\".")
            continue
        cpf = RE_CPF.search(linha)
        doc.itens.append(Item(str(len(doc.itens) + 1), quem, valor, descricao=rotulo or "lista",
                              documento=cpf.group(0) if cpf else "", evidencia=linha))
    docs = []
    for rotulo, doc in listas.items():
        doc.hash = sha256(repr((nome, rotulo, [(i.nome, i.valor) for i in doc.itens])).encode()).hexdigest()
        doc.alertas.extend(avisos)
        if doc.total_impresso is None:
            doc.total_impresso = doc.total
        elif doc.total_impresso != doc.total:
            doc.alertas.append("A soma das linhas não bate com o total informado.")
        docs.append(doc)
    if not docs:
        vazio = Documento(nome, sha256((texto or "").encode()).hexdigest(), "lista", 1)
        vazio.alertas.append("Nenhuma linha com nome e valor (ex.: \"Fulano — R$ 150,00\").")
        docs.append(vazio)
    return docs


def ler_planilha(nome: str, dados: bytes) -> list[Documento]:
    ext = nome.lower().rsplit(".", 1)[-1]
    if ext == "csv":
        texto = dados.decode("utf-8-sig", errors="replace")
        try:
            dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t")
        except csv.Error:
            dialeto = csv.excel
        return ler_tabela(nome, list(csv.reader(io.StringIO(texto), dialeto)))
    if ext == "xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(dados), read_only=True, data_only=True)
        docs = []
        for ws in wb.worksheets:
            linhas = [list(r) for _, r in zip(range(LIMITE_LINHAS + 30), ws.iter_rows(values_only=True))]
            if not any(any(c is not None for c in l) for l in linhas):
                continue
            achados = ler_tabela(nome, linhas, ws.title if len(wb.worksheets) > 1 else "")
            docs += [d for d in achados if d.itens] or ([] if len(wb.worksheets) > 1 else achados)
        wb.close()
        if not docs:
            vazio = Documento(nome, sha256(dados).hexdigest(), "lista", 1)
            vazio.alertas.append("Nenhuma aba com colunas de nome e valor.")
            docs = [vazio]
        return docs
    raise ValueError("planilha .xls antiga não é lida: salve como .xlsx")
